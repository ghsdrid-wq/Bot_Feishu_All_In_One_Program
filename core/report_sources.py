"""Read required local/SMB source sheets once and retain compact counters."""

from __future__ import annotations

import base64
import io
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from .report_formulas import UnsupportedFormula


class SourceUnavailable(RuntimeError):
    pass


def query_text(path: str) -> str:
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("customXml/") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            if root.tag.split("}")[-1] == "DataMashup":
                blob = base64.b64decode(root.text)
                size = int.from_bytes(blob[4:8], "little")
                with zipfile.ZipFile(io.BytesIO(blob[8 : 8 + size])) as package:
                    return "\n".join(
                        package.read(n).decode("utf-8-sig")
                        for n in package.namelist()
                        if n.endswith(".m")
                    )
    return ""


def template_paths(path: str) -> dict:
    text = query_text(path)
    pattern = r'shared\s+(?:#"([^"]+)"|([A-Za-z_][\w]*))\s*=(.*?)(?=\nshared\s|\Z)'
    output = {}
    for quoted, plain, formula in re.findall(pattern, text, re.DOTALL):
        match = re.search(r'File\.Contents\("([^"]+)"\)', formula)
        if match:
            output[quoted or plain] = match.group(1)
    return output


def moment(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip().replace("/", "-"))
        except ValueError:
            match = re.match(
                r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2}) (\d{1,2}):(\d{2})(?::(\d{2}))?",
                value.strip(),
            )
            if match:
                return datetime(*[int(v or 0) for v in match.groups()])
    return None


def column(header, aliases, source):
    for alias in aliases:
        if alias in header:
            return header.index(alias)
    raise SourceUnavailable(f"{source}: missing required column {' / '.join(aliases)}")


class Inputs:
    def __init__(
        self,
        paths: dict,
        business_date: date,
        now: datetime,
        keep_running=lambda: True,
        log=print,
        start_hour=12,
    ):
        self.paths = paths
        self.business_date = business_date
        self.now = now
        self.keep_running = keep_running
        self.log = log
        self.start_hour = start_hour
        self.window_start = datetime.combine(
            business_date, datetime.min.time()
        ).replace(hour=start_hour)
        self.counters = {}
        self.hourly = {}
        self.evidence = []

    def check(self):
        if not self.keep_running():
            raise InterruptedError("Report calculation cancelled")

    def read(self, key, sheet_name=None):
        """Yield rows; check stable file size/mtime before accepting its snapshot."""
        self.check()
        if key not in self.paths:
            raise SourceUnavailable(
                f"No input path configured for required source {key}"
            )
        path = Path(self.paths[key])
        try:
            before = path.stat()
        except OSError as error:
            raise SourceUnavailable(
                f"Cannot access required source {key}: {path}: {error}"
            ) from error
        book = openpyxl.load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
        rows_read = 0
        try:
            name = sheet_name or (
                "DATA" if "DATA" in book.sheetnames else book.sheetnames[0]
            )
            if name not in book.sheetnames:
                raise SourceUnavailable(f"{key}: missing worksheet {name}")
            sheet = book[name]
            sheet.reset_dimensions()
            for row in sheet.iter_rows(values_only=True):
                if rows_read % 2048 == 0:
                    self.check()
                rows_read += 1
                yield row
        finally:
            book.close()
        after = path.stat()
        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            raise SourceUnavailable(
                f"{key}: file changed while being read; retry after its writer finishes"
            )
        self.evidence.append(
            {
                "source": key,
                "path": str(path),
                "rows": max(0, rows_read - 1),
                "bytes": before.st_size,
            }
        )
        self.log(f"[SOURCE] {key}: {max(0, rows_read - 1):,} rows read")

    def autopacking(self, key: str):
        if key in self.hourly:
            return self.hourly[key]
        hour = (int(key) - 1) % 24
        day = self.business_date + timedelta(days=hour < 12)
        slot = datetime.combine(day, datetime.min.time()).replace(hour=hour)
        if slot < self.window_start or slot + timedelta(hours=1) > self.now:
            self.hourly[key] = ({}, {})
            return self.hourly[key]
        path = Path(self.paths.get(key, ""))
        try:
            before = path.stat()
            if datetime.fromtimestamp(before.st_mtime) < slot:
                raise SourceUnavailable(
                    f"{key}: hourly source predates the requested report hour"
                )
        except OSError as error:
            raise SourceUnavailable(
                f"Cannot access required hourly source {key}: {path}: {error}"
            ) from error
        throughput = self.read(key, "Throughput")
        header = [str(x or "").strip() for x in next(throughput, ())]
        pipeline = column(header, ("Pipeline",), key)
        feeder = column(header, ("Feeder No",), key)
        count = column(header, ("Loading Num",), key)
        values = {}
        for row in throughput:
            if not row or row[pipeline] in (None, ""):
                continue
            try:
                index = (str(row[pipeline]).strip(), int(float(row[feeder])))
                values.setdefault(index, float(row[count] or 0))
            except (ValueError, TypeError, IndexError) as error:
                raise SourceUnavailable(f"{key}: invalid feeder/count input") from error
        abnormal = self.read(key, "Abnormal")
        header = [str(x or "").strip() for x in next(abnormal, ())]
        bill = column(header, ("Billcode",), key)
        line = column(header, ("Pipeline",), key)
        errors = {}
        for row in abnormal:
            if len(row) > max(bill, line) and row[bill] not in (None, ""):
                errors.setdefault(str(row[line]), set()).add(str(row[bill]))
        after = path.stat()
        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            raise SourceUnavailable(
                f"{key}: hourly file changed between Throughput and Abnormal reads"
            )
        self.hourly[key] = values, {line: len(bills) for line, bills in errors.items()}
        return self.hourly[key]

    def counter(self, key: str):
        if key in {f"DWS{i}" for i in range(1, 9)} and self.paths.get(key) == "":
            self.log(f"[SOURCE] {key}: unavailable, no local snapshot; no recorded output")
            return Counter()
        if key in self.counters:
            return self.counters[key]
        rows = self.read(key)
        header = [str(x or "").strip() for x in next(rows, ())]
        counter = Counter()
        if key == "REALTIME_DB":
            op = column(
                header,
                ("最新操作类型", "ประเภทการดำเนินงานล่าสุด", "ชื่อประเภทการดำเนินการล่าสุด"),
                key,
            )
            ot = column(header, ("超时类型", "ประเภทเกินเวลา"), key)
            for row in rows:
                if len(row) <= max(op, ot):
                    continue
                hours = re.search(r"\d+", str(row[ot] or ""))
                if row[op] is not None and hours:
                    counter[(str(row[op]).casefold(), int(hours.group()))] += 1
        else:
            scan = key in ("AUTO_PDA_DATA", "DWS_PDA_DATA")
            times = (
                ("扫描时间", "เวลาที่สแกน", "เวลาที่ทำการสแกน", "เวลาสแกน")
                if scan or key == "DWS9-11"
                else ("出秤时间",)
            )
            time_index = column(header, times, key)
            scanner = (
                column(header, ("扫描员", "พนักงานสแกน", "ผู้สแกน"), key)
                if scan
                else column(header, ("dws序号",), key)
                if key == "DWS9-11"
                else None
            )
            bag = (
                column(header, ("所属包号", "หมายเลขกระสอบ"), key)
                if key == "DWS_PDA_DATA"
                else None
            )
            dated_rows = 0
            invalid_times = 0
            for row in rows:
                if len(row) <= max(
                    i for i in (time_index, scanner, bag) if i is not None
                ):
                    continue
                if bag is not None and str(row[bag] or "").strip():
                    continue
                timestamp = moment(row[time_index])
                if timestamp is None:
                    if row[time_index] not in (None, ""):
                        invalid_times += 1
                    continue
                dated_rows += 1
                if timestamp.date() not in (
                    self.business_date,
                    date.fromordinal(self.business_date.toordinal() + 1),
                ):
                    continue
                if not self.window_start <= timestamp <= self.now:
                    continue
                person = (
                    str(row[scanner] or "").strip().casefold()
                    if scanner is not None
                    else ""
                )
                counter[(timestamp.date(), timestamp.hour, person)] += 1
            if dated_rows and not counter and key not in {f"DWS{i}" for i in range(1, 9)}:
                raise SourceUnavailable(
                    f"{key}: non-empty source has no rows in the requested report window"
                )
            if invalid_times:
                raise SourceUnavailable(
                    f"{key}: {invalid_times} non-empty scan timestamps could not be parsed"
                )
        self.counters[key] = counter
        return counter

    def count(self, refs, criteria):
        fields = []
        keys = []
        for ref in refs:
            if ref.sheet == "@table":
                match = re.fullmatch(r"([^\[]+)\[([^\]]+)\]", ref.address)
                if not match:
                    raise UnsupportedFormula(ref.address)
                key, field = match.groups()
                field = {"SCANNER": "person", "SCAN_TIME": "time"}.get(field)
            else:
                key, field = ref.sheet, ref.address
                if key == "REALTIME DB DATA":
                    key, field = (
                        "REALTIME_DB",
                        {"A:A": "operation", "C:C": "overtime"}.get(field),
                    )
                elif key == "DWS9-11":
                    field = {"A:A": "time", "E:E": "person"}.get(field)
                elif re.fullmatch(r"DWS[1-8]", key):
                    field = "time" if field == "H:H" else None
                else:
                    raise UnsupportedFormula(f"Unknown source range {ref}")
            if field is None:
                raise UnsupportedFormula(f"Unknown source column {ref}")
            fields.append(field)
            keys.append(key)
        if len(set(keys)) != 1:
            raise UnsupportedFormula("COUNTIFS spans different source files")
        data = self.counter(keys[0])
        if keys[0] == "REALTIME_DB":
            filters = dict(zip(fields, criteria))
            return sum(
                count
                for (op, hours), count in data.items()
                if op == str(filters.get("operation", op)).casefold()
                and hours == filters.get("overtime", hours)
            )
        time_filter = None
        person = None
        for field, criterion in zip(fields, criteria):
            if field == "person":
                person = str(criterion).casefold()
            else:
                match = re.fullmatch(
                    r"(\d{4})[/-](\d{1,2})[/-](\d{1,2}) (\d{1,2})\*", str(criterion)
                )
                if not match:
                    raise UnsupportedFormula(
                        f"Unsupported time criterion {criterion!r}"
                    )
                year, month, day, hour = map(int, match.groups())
                time_filter = (date(year, month, day), hour)
        return sum(
            count
            for (day, hour, scanner), count in data.items()
            if (time_filter is None or (day, hour) == time_filter)
            and (person is None or person == scanner)
        )
