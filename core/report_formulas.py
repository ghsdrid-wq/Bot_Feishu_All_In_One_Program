"""Evaluate the small, audited formula vocabulary of the report templates.

This is not a general Excel engine. Unsupported formulas fail explicitly;
Excel itself is never started, and no Python eval/exec is used.
"""

from __future__ import annotations

import math
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from typing import ClassVar

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel, to_excel


class UnsupportedFormula(ValueError):
    pass


class FormulaError(ValueError):
    pass


@dataclass(frozen=True)
class Reference:
    sheet: str
    address: str


class Parser:
    PRECEDENCE: ClassVar[dict[str, int]] = {
        "=": 10,
        "<>": 10,
        "<": 10,
        ">": 10,
        "<=": 10,
        ">=": 10,
        "&": 20,
        "+": 30,
        "-": 30,
        "*": 40,
        "/": 40,
        "^": 50,
    }

    def __init__(self, formula: str):
        self.tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
        self.index = 0

    def take(self):
        token = self.tokens[self.index]
        self.index += 1
        return token

    def expression(self, minimum=0):
        token = self.take()
        if token.type == "OPERATOR-PREFIX":
            left = ("unary", token.value, self.expression(60))
        elif token.type == "FUNC" and token.subtype == "OPEN":
            args = []
            while (
                self.index < len(self.tokens)
                and self.tokens[self.index].subtype != "CLOSE"
            ):
                args.append(self.expression())
                if self.tokens[self.index].type != "SEP":
                    break
                self.take()
            if self.take().subtype != "CLOSE":
                raise UnsupportedFormula("Unclosed function")
            left = ("call", token.value[:-1].upper(), args)
        elif token.type == "PAREN" and token.subtype == "OPEN":
            left = self.expression()
            if self.take().subtype != "CLOSE":
                raise UnsupportedFormula("Unclosed expression")
        elif token.subtype == "NUMBER":
            left = ("value", float(token.value))
        elif token.subtype == "TEXT":
            left = ("value", token.value[1:-1].replace('""', '"'))
        elif token.subtype == "LOGICAL":
            left = ("value", token.value == "TRUE")
        elif token.subtype == "RANGE":
            left = ("reference", token.value)
        elif token.subtype == "ERROR":
            left = ("error", token.value)
        else:
            raise UnsupportedFormula(f"Unsupported token {token.value!r}")
        while self.index < len(self.tokens):
            token = self.tokens[self.index]
            priority = (
                self.PRECEDENCE.get(token.value, -1)
                if token.type == "OPERATOR-INFIX"
                else -1
            )
            if priority < minimum:
                break
            self.take()
            right = self.expression(priority + (0 if token.value == "^" else 1))
            left = ("binary", token.value, left, right)
        return left

    def parse(self):
        tree = self.expression()
        if self.index != len(self.tokens):
            raise UnsupportedFormula("Unexpected trailing formula tokens")
        return tree


def matches(value, criterion) -> bool:
    if not isinstance(criterion, str):
        return value == criterion
    match = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", criterion)
    operator, expected = match.groups()
    operator = operator or "="
    try:
        number = float(expected)
        if not isinstance(value, (int, float)):
            return False
        return {
            "=": value == number,
            "<>": value != number,
            ">": value > number,
            "<": value < number,
            ">=": value >= number,
            "<=": value <= number,
        }[operator]
    except ValueError:
        actual = str(value or "")
        pattern = "".join(
            ".*" if c == "*" else "." if c == "?" else re.escape(c) for c in expected
        )
        equal = re.fullmatch(pattern, actual, re.IGNORECASE) is not None
        if operator not in ("=", "<>"):
            raise UnsupportedFormula("Unsupported text comparison criterion")
        return not equal if operator == "<>" else equal


class Calculator:
    def __init__(self, workbook, now: datetime, source_count: Callable, overrides=None):
        self.book = workbook
        self.now = now
        self.source_count = source_count
        self.values = dict(overrides or {})
        self.active = set()
        self.parsed = {}

    def reference(self, sheet: str, text: str) -> Reference:
        if "[" in text:
            return Reference("@table", text)
        if "!" in text:
            name, text = text.rsplit("!", 1)
            sheet = name.strip("'").replace("''", "'")
        return Reference(sheet, text.replace("$", ""))

    def scalar(self, value):
        if isinstance(value, Reference):
            if ":" in value.address or value.sheet == "@table":
                raise UnsupportedFormula("Expected a single cell")
            return self.cell(value.sheet, value.address)
        return value

    def numeric(self, value):
        value = self.scalar(value)
        return 0 if value in (None, "") else float(value)

    def flatten(self, value):
        if isinstance(value, Reference):
            if value.sheet == "@table":
                raise UnsupportedFormula("Unbounded table cannot be materialized")
            first_col, first_row, last_col, last_row = range_boundaries(value.address)
            if first_row is None or last_row is None:
                raise UnsupportedFormula("Unbounded column cannot be materialized")
            from openpyxl.utils.cell import get_column_letter

            return [
                self.cell(value.sheet, f"{get_column_letter(c)}{r}")
                for r in range(first_row, last_row + 1)
                for c in range(first_col, last_col + 1)
            ]
        return [self.scalar(value)]

    def cell(self, sheet: str, address: str):
        key = (sheet, address)
        if key in self.values:
            return self.values[key]
        if key in self.active:
            raise UnsupportedFormula(f"Circular report formula: {sheet}!{address}")
        cell = self.book[sheet][address]
        value = cell.value
        if cell.data_type == "f":
            if not isinstance(value, str):
                raise UnsupportedFormula(f"Unmapped array formula: {sheet}!{address}")
            self.active.add(key)
            try:
                if value not in self.parsed:
                    self.parsed[value] = Parser(value).parse()
                value = self.evaluate(sheet, self.parsed[value])
                value = self.scalar(value)
            except (ZeroDivisionError, OverflowError) as error:
                raise FormulaError(str(error)) from error
            finally:
                self.active.remove(key)
        elif isinstance(value, datetime):
            value = to_excel(value)
        elif cell.data_type == "e":
            raise FormulaError(str(value))
        self.values[key] = value
        return value

    def evaluate(self, sheet: str, node):
        kind = node[0]
        if kind == "value":
            return node[1]
        if kind == "error":
            raise FormulaError(node[1])
        if kind == "reference":
            return self.reference(sheet, node[1])
        if kind == "unary":
            value = self.numeric(self.evaluate(sheet, node[2]))
            return -value if node[1] == "-" else value
        if kind == "binary":
            left, right = (self.scalar(self.evaluate(sheet, n)) for n in node[2:])
            operator = node[1]
            if operator == "&":

                def text(v):
                    return (
                        str(int(v))
                        if isinstance(v, float) and v.is_integer()
                        else ""
                        if v is None
                        else str(v)
                    )

                return text(left) + text(right)
            if operator in ("=", "<>", "<", ">", "<=", ">="):
                left, right = (0 if v is None else v for v in (left, right))
                if isinstance(left, str) and isinstance(right, str):
                    left, right = left.casefold(), right.casefold()
                if operator == "=":
                    return left == right
                if operator == "<>":
                    return left != right
                if operator == "<":
                    return left < right
                if operator == ">":
                    return left > right
                if operator == "<=":
                    return left <= right
                return left >= right
            a, b = self.numeric(left), self.numeric(right)
            if operator == "+":
                return a + b
            if operator == "-":
                return a - b
            if operator == "*":
                return a * b
            if operator == "/":
                if b == 0:
                    raise FormulaError("Division by zero")
                return a / b
            if operator == "^":
                return a**b
            raise UnsupportedFormula(operator)
        name, args = node[1:]
        if name == "IF":
            condition = self.scalar(self.evaluate(sheet, args[0]))
            return (
                self.evaluate(sheet, args[1] if condition else args[2])
                if condition or len(args) > 2
                else False
            )
        if name == "IFERROR":
            try:
                return self.scalar(self.evaluate(sheet, args[0]))
            except FormulaError:
                return self.evaluate(sheet, args[1])
        values = [self.evaluate(sheet, arg) for arg in args]
        if name == "NOW":
            return to_excel(self.now)
        if name == "TODAY":
            return to_excel(self.now.replace(hour=0, minute=0, second=0, microsecond=0))
        if name == "HOUR":
            return from_excel(self.numeric(values[0])).hour
        if name == "TEXT":
            date = from_excel(self.numeric(values[0]))
            fmt = self.scalar(values[1])
            if fmt == "yyyy-mm-dd":
                return date.strftime("%Y-%m-%d")
            if fmt == "yyyy/mm/dd":
                return date.strftime("%Y/%m/%d")
            if fmt == "yyyy/m/dd":
                return f"{date.year}/{date.month}/{date.day:02d}"
            if fmt == "yyyy/m/d":
                return f"{date.year}/{date.month}/{date.day}"
            raise UnsupportedFormula(f"Unsupported date format {fmt}")
        if name in ("SUM", "MAX", "COUNT"):
            numbers = [
                x
                for v in values
                for x in self.flatten(v)
                if isinstance(x, (int, float)) and not isinstance(x, bool)
            ]
            return (
                sum(numbers)
                if name == "SUM"
                else len(numbers)
                if name == "COUNT"
                else max(numbers, default=0)
            )
        if name in ("COUNTIF", "COUNTIFS"):
            refs, criteria = values[::2], [self.scalar(v) for v in values[1::2]]
            if any(
                r.sheet == "@table" or re.fullmatch(r"[A-Z]+:[A-Z]+", r.address)
                for r in refs
            ):
                return self.source_count(refs, criteria)
            columns = [self.flatten(r) for r in refs]
            return sum(
                all(matches(v, c) for v, c in zip(row, criteria))
                for row in zip(*columns)
            )
        if name in ("SUMIF", "AVERAGEIF"):
            criteria_values = self.flatten(values[0])
            sum_values = self.flatten(values[2] if len(values) > 2 else values[0])
            chosen = [
                v
                for match, v in zip(criteria_values, sum_values)
                if matches(match, self.scalar(values[1]))
                and isinstance(v, (int, float))
            ]
            if name == "SUMIF":
                return sum(chosen)
            if not chosen:
                raise FormulaError("No matching average values")
            return sum(chosen) / len(chosen)
        if name == "INT":
            return math.floor(self.numeric(values[0]))
        if name == "VLOOKUP":
            key, ref, column = (
                self.scalar(values[0]),
                values[1],
                int(self.numeric(values[2])),
            )
            if len(values) > 3 and self.scalar(values[3]) not in (0, False):
                raise UnsupportedFormula("Approximate VLOOKUP is not supported")
            first_col, first_row, _last_col, last_row = range_boundaries(ref.address)
            from openpyxl.utils.cell import get_column_letter

            for row in range(first_row, last_row + 1):
                if self.cell(ref.sheet, f"{get_column_letter(first_col)}{row}") == key:
                    return self.cell(
                        ref.sheet, f"{get_column_letter(first_col + column - 1)}{row}"
                    )
            raise FormulaError("VLOOKUP key was not found")
        raise UnsupportedFormula(f"Unsupported function {name}")

    def calculate(self):
        for sheet in self.book:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        self.cell(sheet.title, cell.coordinate)
        return self.values
