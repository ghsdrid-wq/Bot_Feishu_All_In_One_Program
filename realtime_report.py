# -*- coding: utf-8 -*-
"""แปลงไฟล์ดิบ "ควบคุมติดตามแบบเรียลไทม์DB" จาก JMS เป็นไฟล์รายงานพร้อมส่ง

ของเดิมงานนี้ทำใน Excel ด้วย Power Query + COUNTIFS ในไฟล์ 1DWS_&_PDA ซึ่ง
ผูกพัสดุเกินเวลา (งานของ QC) ไว้กับยอด KPI ทั้งที่เป็นคนละงานกัน ย้ายมาทำใน
โปรแกรมแทน ขั้นตอน Realtime DB จึงจบในตัวเอง คือดึงจาก JMS แล้วได้ไฟล์ที่มี
ทั้งชีตตารางสรุปและชีตข้อมูลดิบในไฟล์เดียว

ตารางที่ได้หน้าตาเหมือนของเดิมทุกอย่าง เพราะฝ่าย QC ดูตัวนี้อยู่แล้ว
"""

import os
import re
import shutil
from datetime import datetime
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LogFunc = Callable[..., None]

SHEET_TABLE = "TABLE"
SHEET_RAW = "DATA"

# หัวคอลัมน์ในไฟล์ดิบที่ต้องใช้ ชื่อมาจาก JMS ตรง ๆ
COL_OPERATE = "ประเภทการดำเนินงานล่าสุด"
COL_OVERTIME = "ประเภทเกินเวลา"

# แถวของตาราง (ชื่อที่โชว์, ค่าที่ JMS ส่งมา) เรียงตามของเดิมใน Excel
# สามแถวท้ายยังไม่ได้ขอจาก JMS ในตอนนี้ แต่คงไว้ให้หน้าตาตารางเหมือนเดิม
OPERATE_ROWS: Sequence[Tuple[str, str]] = (
    ("Arrive Scan", "arrive scan"),
    ("Packing Scan", "packing scan"),
    ("Problematic Parcel Scan", "problematic parcel scan"),
    ("Return Item Registration", "return item registration"),
    ("Sending Scan", "sending scan"),
    ("DP Departure Scan", "dp departure scan"),
    ("DC Departure Scan", "dc departure scan"),
)

# คอลัมน์ชั่วโมงที่เกิน ตรงกับที่ขอไปใน payload ของ JMS
OVERTIME_HOURS: Sequence[int] = (4, 6, 8, 12, 24, 36, 48)

TITLE = "พัสดุเกินเวลา 4-48  ชั่วโมง"

# สีเดียวกับไฟล์เดิม (Office accent1 อ่อน 40% และ 80%) เขียนเป็นรหัสสีตรง ๆ
# จะได้ไม่ต้องพึ่งธีมของเครื่องที่เปิดไฟล์
HEAD_FILL = PatternFill("solid", fgColor="8EAADB")
SUM_FILL = PatternFill("solid", fgColor="D9E2F3")
FONT_NAME = "Microsoft YaHei"
ROW_HEIGHT = 20.1
TITLE_HEIGHT = 53.45


def _log(log: Optional[LogFunc], msg: str, level: Optional[str] = None) -> None:
    if log is None:
        print(msg)
        return
    try:
        log(msg, level) if level else log(msg)
    except TypeError:
        log(msg)


def _overtime_hours(value) -> Optional[int]:
    """ดึงตัวเลขชั่วโมงออกจากข้อความอย่าง "Exceed 12 hours with no track"

    อ่านด้วย regex แทนการเทียบข้อความเต็ม เผื่อ JMS แก้ถ้อยคำภายหลัง
    """
    if value is None:
        return None
    match = re.search(r"(\d+)", str(value))
    return int(match.group(1)) if match else None


def read_raw(path: str) -> Tuple[List[str], List[tuple]]:
    """อ่านไฟล์ดิบทั้งไฟล์ คืนหัวคอลัมน์กับข้อมูล"""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        # ปกติไฟล์ดิบมีชีตเดียว แต่ถ้าเผลอชี้มาที่ไฟล์ที่แปลงแล้ว ให้หยิบชีต
        # ข้อมูลดิบในนั้นแทน จะได้ไม่ไปอ่านตารางสรุปมาเป็นข้อมูลตั้งต้น
        names = {n.casefold(): n for n in book.sheetnames}
        name = names.get(SHEET_RAW.casefold(), book.sheetnames[0])
        sheet = book[name]
        # ไฟล์ที่ JMS ส่งมาประกาศขนาดตารางไว้ผิด (บอกว่ามีคอลัมน์เดียว) โหมด
        # read_only เชื่อค่านั้นแล้วตัดคอลัมน์ที่เหลือทิ้ง ต้องสั่งให้อ่านของจริง
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = [str(x).strip() if x is not None else "" for x in next(rows, ())]
        width = len(header)
        data = []
        for row in rows:
            if not any(x is not None and str(x).strip() for x in row):
                continue
            # แถวที่สั้นกว่าหัวตารางเกิดได้เมื่อเซลล์ท้ายแถวว่าง เติมให้เท่ากัน
            data.append(tuple(row[:width]) + (None,) * max(0, width - len(row)))
        return header, data
    finally:
        book.close()


def build_matrix(header: Sequence[str], data: Sequence[tuple],
                 log: Optional[LogFunc] = None) -> Dict[str, Dict[int, int]]:
    """นับจำนวนพัสดุแยกตามประเภทการดำเนินงาน x ชั่วโมงที่เกิน"""
    try:
        i_operate = header.index(COL_OPERATE)
        i_overtime = header.index(COL_OVERTIME)
    except ValueError as exc:
        raise Exception(
            f"ไฟล์ดิบไม่มีคอลัมน์ที่ต้องใช้ ({COL_OPERATE} / {COL_OVERTIME})"
        ) from exc

    lookup = {source: label for label, source in OPERATE_ROWS}
    counts: Dict[str, Dict[int, int]] = {
        label: {hour: 0 for hour in OVERTIME_HOURS} for label, _ in OPERATE_ROWS
    }
    unmatched: Dict[str, int] = {}

    for row in data:
        operate = str(row[i_operate] or "").strip().lower()
        hours = _overtime_hours(row[i_overtime])
        label = lookup.get(operate)
        if label is None or hours not in counts[label]:
            key = f"{row[i_operate]} / {row[i_overtime]}"
            unmatched[key] = unmatched.get(key, 0) + 1
            continue
        counts[label][hours] += 1

    # ถ้ามีค่าที่ไม่เข้าช่องไหนเลย ต้องบอก ไม่ใช่ปล่อยให้ยอดหายเงียบ ๆ
    for key, count in sorted(unmatched.items(), key=lambda kv: -kv[1]):
        _log(log, f"Realtime DB — ไม่รู้จักค่า {key} ({count} ชิ้น) ไม่ถูกนับในตาราง", "WARN")
    return counts


def _write_table(sheet, counts: Dict[str, Dict[int, int]], pulled_at: datetime) -> None:
    columns = len(OVERTIME_HOURS) + 2  # ชื่อแถว + ชั่วโมง + ผลรวม

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    title = sheet.cell(row=1, column=1, value=f"{TITLE} {pulled_at:%Y-%m-%d}")
    title.font = Font(name=FONT_NAME, size=18)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = TITLE_HEIGHT

    head = [COL_OPERATE] + [f"{h} ชั่วโมง" for h in OVERTIME_HOURS] + ["ผลรวม"]
    for col, text in enumerate(head, start=1):
        cell = sheet.cell(row=2, column=col, value=text)
        cell.fill = HEAD_FILL
        cell.font = Font(name=FONT_NAME, size=14)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    totals = {hour: 0 for hour in OVERTIME_HOURS}
    for offset, (label, _source) in enumerate(OPERATE_ROWS):
        row = 3 + offset
        cell = sheet.cell(row=row, column=1, value=label)
        cell.font = Font(name=FONT_NAME, size=14)
        line = 0
        for col, hour in enumerate(OVERTIME_HOURS, start=2):
            value = counts[label][hour]
            totals[hour] += value
            line += value
            number = sheet.cell(row=row, column=col, value=value)
            number.font = Font(name=FONT_NAME, size=14)
        last = sheet.cell(row=row, column=columns, value=line)
        last.font = Font(name=FONT_NAME, size=14)

    sum_row = 3 + len(OPERATE_ROWS)
    label = sheet.cell(row=sum_row, column=1, value="ผลรวม")
    label.fill = SUM_FILL
    label.font = Font(name=FONT_NAME, size=14)
    label.alignment = Alignment(horizontal="center", vertical="center")
    for col, hour in enumerate(OVERTIME_HOURS, start=2):
        cell = sheet.cell(row=sum_row, column=col, value=totals[hour])
        cell.fill = SUM_FILL
        cell.font = Font(name=FONT_NAME, size=14)
    grand = sheet.cell(row=sum_row, column=columns, value=sum(totals.values()))
    grand.fill = SUM_FILL
    grand.font = Font(name=FONT_NAME, size=14)

    for row in range(2, sum_row + 1):
        sheet.row_dimensions[row].height = ROW_HEIGHT
    sheet.column_dimensions["A"].width = 32.9
    for col in range(2, columns + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 15.7


def _write_raw(sheet, header: Sequence[str], data: Sequence[tuple]) -> None:
    sheet.append(list(header))
    for row in data:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    # ตั้งความกว้างเองตรงนี้ จะได้ไม่ต้องให้ autofit ของโปรแกรมไปรื้อความกว้าง
    # ที่จัดไว้ในชีตตารางสรุป วัดจากพันแถวแรกก็พอ ไฟล์ใหญ่จะได้ไม่ช้า
    widths = [len(str(x)) for x in header]
    for row in data[:1000]:
        for i, value in enumerate(row):
            if value is None or i >= len(widths):
                continue
            widths[i] = max(widths[i], len(str(value)))
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)


def table_range() -> str:
    """ช่วงเซลล์ของตาราง ใช้ตั้งค่า Export รูปในหน้าไฟล์ Excel"""
    columns = len(OVERTIME_HOURS) + 2
    return f"A1:{get_column_letter(columns)}{2 + len(OPERATE_ROWS) + 1}"


def build_report(raw_path: str, log: Optional[LogFunc] = None,
                 pulled_at: Optional[datetime] = None) -> str:
    """แปลงไฟล์ดิบที่ดาวน์โหลดมา ให้กลายเป็นไฟล์รายงานที่ตำแหน่งเดิม

    เขียนลงไฟล์ชั่วคราวก่อนแล้วค่อยสลับทับ ถ้าพังกลางทางไฟล์เดิมจะยังอยู่
    ไม่เหลือไฟล์ครึ่ง ๆ กลาง ๆ ให้ส่งเข้ากลุ่ม
    """
    pulled_at = pulled_at or datetime.now()
    header, data = read_raw(raw_path)
    counts = build_matrix(header, data, log=log)

    book = openpyxl.Workbook()
    table = book.active
    table.title = SHEET_TABLE
    _write_table(table, counts, pulled_at)
    _write_raw(book.create_sheet(SHEET_RAW), header, data)

    tmp_path = f"{raw_path}.tmp"
    book.save(tmp_path)
    book.close()
    os.replace(tmp_path, raw_path)

    total = sum(sum(v.values()) for v in counts.values())
    _log(log, f"Realtime DB — สร้างตารางแล้ว {total:,} ชิ้น จากข้อมูลดิบ {len(data):,} แถว")
    return raw_path


if __name__ == "__main__":
    import sys
    source = sys.argv[1]
    target = sys.argv[2] if len(sys.argv) > 2 else None
    if target:
        shutil.copyfile(source, target)
        source = target
    build_report(source)
