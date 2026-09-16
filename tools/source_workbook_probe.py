from __future__ import annotations

import argparse
import json
import re
import sys
import time
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def count_xml_rows(path: Path) -> dict[str, int]:
    counts: dict[str, int] = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                continue
            row_count = 0
            carry = b""
            with archive.open(name) as stream:
                while chunk := stream.read(4 * 1024 * 1024):
                    data = carry + chunk
                    row_count += data.count(b"<row")
                    carry = data[-3:]
            counts[name] = row_count
    return counts


def probe(path: Path) -> dict:
    started = time.perf_counter()
    xml_rows = count_xml_rows(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            if sheet.max_row == 1 and xml_rows.get(f"xl/worksheets/sheet{index}.xml", 0) > 1:
                sheet.reset_dimensions()
            first_rows = []
            for row in sheet.iter_rows(min_row=1, max_row=3, values_only=True):
                first_rows.append([str(value)[:80] if value is not None else None for value in row[:20]])
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": xml_rows.get(f"xl/worksheets/sheet{index}.xml", sheet.max_row),
                    "columns": max((len(row) for row in first_rows), default=0),
                    "first_rows": first_rows,
                }
            )
        return {
            "path": str(path.resolve()),
            "size_bytes": path.stat().st_size,
            "seconds": round(time.perf_counter() - started, 3),
            "sheets": sheets,
        }
    finally:
        workbook.close()


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+", type=Path)
    args = parser.parse_args()
    print(json.dumps([probe(path) for path in args.paths], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
