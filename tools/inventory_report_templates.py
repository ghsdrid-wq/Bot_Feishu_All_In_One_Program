"""Inventory all workbook formulas and Power Query dependencies read-only."""

from __future__ import annotations

import base64
import io
import json
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def queries(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("customXml/") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            if root.tag.split("}")[-1] != "DataMashup":
                continue
            blob = base64.b64decode(root.text)
            length = int.from_bytes(blob[4:8], "little")
            with zipfile.ZipFile(io.BytesIO(blob[8 : 8 + length])) as package:
                return "\n".join(
                    package.read(n).decode("utf-8-sig")
                    for n in package.namelist()
                    if n.endswith(".m")
                )
    return ""


def inventory(path: Path, folder: Path) -> dict:
    book = openpyxl.load_workbook(
        path, read_only=False, data_only=False, keep_links=True
    )
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {
        "file": path.name,
        "sheets": [],
        "defined_names": list(book.defined_names),
        "sources": [],
    }
    for sheet in book:
        formulas = {
            cell.coordinate: cell.value
            for row in sheet
            for cell in row
            if cell.data_type == "f"
        }
        kinds = Counter(
            re.match(r"=?([A-Z_.]+)", str(value)).group(1)
            if re.match(r"=?([A-Z_.]+)", str(value))
            else "other"
            for value in formulas.values()
        )
        result["sheets"].append(
            {
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "state": sheet.sheet_state,
                "formula_count": len(formulas),
                "formula_kinds": dict(kinds),
                "formulas": formulas,
                "tables": list(sheet.tables),
                "merged": [str(r) for r in sheet.merged_cells.ranges],
                "sample_formula_cache": {
                    addr: cached[sheet.title][addr].value for addr in list(formulas)[:4]
                },
            }
        )
    text = queries(path)
    result["sources"] = re.findall(r'File\.Contents\("([^"]+)"\)', text)
    result["query_names"] = re.findall(r"shared\s+(.*?)\s*=", text)
    (folder / (path.stem + ".pq")).write_text(text, encoding="utf-8")
    book.close()
    cached.close()
    return result


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    folder = ROOT / "out/workbook_migration"
    folder.mkdir(parents=True, exist_ok=True)
    results = [
        inventory(ROOT / "Excel" / name, folder)
        for name in ("1AUTOPACKING_BOT.xlsx", "1DWS_&_PDA_v1.8.3 Bot.xlsx")
    ]
    (folder / "inventory.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    for result in results:
        print(
            result["file"],
            "queries",
            len(result["query_names"]),
            "sources",
            len(result["sources"]),
        )
        for sheet in result["sheets"]:
            print(
                json.dumps(
                    {
                        k: sheet[k]
                        for k in (
                            "name",
                            "rows",
                            "columns",
                            "formula_count",
                            "formula_kinds",
                            "tables",
                        )
                    },
                    ensure_ascii=False,
                )
            )


if __name__ == "__main__":
    main()
