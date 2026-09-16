import configparser
import gc
import sys
import tempfile
import time
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

import Createphoto  # noqa: E402


def build_test_config(workbook_path: Path) -> configparser.ConfigParser:
    config = configparser.ConfigParser()
    config["TIME"] = {"start_hour": "12", "end_hour": "12", "run_minute": "5"}
    config["WORKBOOKS"] = {"items": "TEST"}
    config["WORKBOOK:TEST"] = {
        "path": str(workbook_path),
        "display_name": "Excel cleanup test",
        "enabled": "true",
    }
    config["EXPORTS"] = {"items": "TEST"}
    config["EXPORT:TEST"] = {
        "enabled": "true",
        "send_enabled": "false",
        "workbook": "TEST",
        "sheet": "Sheet1",
        "range": "A1:C3",
        "file": "cleanup-test.png",
        "delete_by_start": "false",
    }
    return config


def create_test_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    colors = ["1F4E78", "70AD47", "FFC000"]
    for row in range(1, 4):
        for column in range(1, 4):
            cell = worksheet.cell(row=row, column=column, value=f"R{row}C{column}")
            cell.fill = PatternFill("solid", fgColor=colors[column - 1])
    workbook.save(path)
    workbook.close()


def main() -> int:
    logs = []
    captured_pids = []
    original_load_config = Createphoto.load_config
    original_save_config = Createphoto.save_config
    original_get_excel_pid = Createphoto.get_excel_pid

    with tempfile.TemporaryDirectory(prefix="autoreport-cleanup-") as temp_dir:
        root = Path(temp_dir)
        workbook_path = root / "cleanup-test.xlsx"
        output_dir = root / "output"
        create_test_workbook(workbook_path)
        config = build_test_config(workbook_path)

        def capture_pid(excel):
            pid = original_get_excel_pid(excel)
            captured_pids.append(pid)
            return pid

        Createphoto.load_config = lambda: config
        Createphoto.save_config = lambda _config: None
        Createphoto.get_excel_pid = capture_pid
        pipeline_error = None
        try:
            Createphoto.run_create(str(output_dir), log=logs.append)
        except Exception as error:
            pipeline_error = error
        finally:
            Createphoto.load_config = original_load_config
            Createphoto.save_config = original_save_config
            Createphoto.get_excel_pid = original_get_excel_pid
            gc.collect()

        image_path = output_dir / "cleanup-test.png"
        pid = captured_pids[-1] if captured_pids else None
        time.sleep(1)
        process_leftover = Createphoto.is_process_alive(pid)

        print(f"PIPELINE_COMPLETED={pipeline_error is None}")
        if pipeline_error is not None:
            print(f"PIPELINE_ERROR={pipeline_error}")
        print(f"IMAGE_EXISTS={image_path.exists()}")
        print(f"IMAGE_SIZE={image_path.stat().st_size if image_path.exists() else 0}")
        print(f"EXCEL_PID={pid}")
        print(f"EXCEL_LEFTOVER={process_leftover}")
        for entry in logs:
            if "CLEANUP" in entry or "WARN" in entry:
                print(f"LOG={entry}")

        return 0 if not process_leftover else 1


if __name__ == "__main__":
    raise SystemExit(main())
