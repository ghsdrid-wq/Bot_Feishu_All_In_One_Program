"""Behavioral checks for Python-calculated, values-only template copies."""

from __future__ import annotations

import configparser
import hashlib
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.report_formulas import Calculator, UnsupportedFormula
from core.report_sources import Inputs, SourceUnavailable
from report_workbooks import calculate, prepare


def fixture(path, headers, rows, sheet="Sheet1", extra=None):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if extra:
        for name, (head, data) in extra.items():
            ws = book.create_sheet(name)
            ws.append(head)
            for row in data:
                ws.append(row)
    book.save(path)
    book.close()


def input_fixtures(folder):
    auto = folder / "hour.xlsx"
    fixture(
        auto,
        ["Pipeline", "Feeder No", "Loading Num"],
        [["58987-001", 1, 201], ["58987-001", 1, 999], ["58987-002", 1, 199]],
        "Throughput",
        {
            "Abnormal": (
                ["Billcode", "Pipeline"],
                [
                    ["A1", "58987-001"],
                    ["A1", "58987-001"],
                    ["A2", "58987-001"],
                    ["B1", "58987-002"],
                ],
            )
        },
    )
    dws = folder / "DWS1.xlsx"
    fixture(
        dws, ["单号", "出秤时间"], [[str(i), "2026/10/20 15:05:00"] for i in range(201)]
    )
    dws911 = folder / "DWS9-11.xlsx"
    fixture(
        dws911,
        ["扫描时间", "dws序号"],
        [["2026-10-20 15:05:00", "10.30.32.220"] for _ in range(202)],
    )
    scan = folder / "scan.xlsx"
    fixture(
        scan,
        ["运单号", "所属包号", "扫描时间", "扫描员", "上传时间"],
        [
            [str(i), "", "2026-10-20 15:05:00", "Autopacking_pda01", ""]
            for i in range(203)
        ]
        + [[str(i), "", "2026-10-20 15:05:00", "Dws_pda01", ""] for i in range(204)]
        + [["bagged", "B1", "2026-10-20 15:05:00", "Dws_pda01", ""]],
    )
    realtime = folder / "realtime.xlsx"
    fixture(
        realtime,
        ["最新操作类型", "超时类型"],
        [
            ["Arrive scan", "Exceed 4 hours with no track"],
            ["Packing scan", "Exceed 12 hours with no track"],
        ],
    )
    return {
        **{f"{h:02d}": str(auto) for h in range(24)},
        **{f"DWS{i}": str(dws) for i in range(1, 9)},
        "DWS9-11": str(dws911),
        "AUTO_PDA_DATA": str(scan),
        "DWS_PDA_DATA": str(scan),
        "REALTIME_DB": str(realtime),
    }


class FormulaTests(unittest.TestCase):
    def test_lazy_branches_threshold_and_unsupported_function(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Test"
        for address, value in {
            "A1": 200,
            "A2": 201,
            "A3": 0,
            "B1": '=SUMIF(A1:A3,">200",A1:A3)',
            "B2": "=IFERROR(1/0,17)",
            "B3": "=IF(0,1/0,SUM(A1:A3))",
            "B4": "=INT(9/2)",
        }.items():
            sheet[address] = value
        calc = Calculator(book, datetime(2026, 10, 21), lambda *_: 0)
        self.assertEqual(
            [calc.cell("Test", f"B{i}") for i in range(1, 5)], [201, 17, 401, 4]
        )
        sheet["B6"] = '=HOUR(NOW())&":00"'
        self.assertEqual(calc.cell("Test", "B6"), "0:00")
        sheet["B5"] = "=IFERROR(UNSUPPORTED(A1),0)"
        with self.assertRaises(UnsupportedFormula):
            calc.cell("Test", "B5")

    def test_missing_file_is_not_a_zero_count(self):
        with tempfile.TemporaryDirectory() as tmp:
            inputs = Inputs(
                {"DWS1": os.path.join(tmp, "missing.xlsx")},
                date(2026, 10, 20),
                datetime(2026, 10, 21),
            )
            with self.assertRaises(SourceUnavailable):
                inputs.counter("DWS1")

    def test_future_hour_does_not_require_a_missing_shared_file(self):
        inputs = Inputs(
            {}, date(2026, 10, 20), datetime(2026, 10, 21, 0, 5), start_hour=15
        )
        self.assertEqual(inputs.autopacking("03"), ({}, {}))

    def test_thai_chinese_aliases_and_start_hour_match(self):
        with tempfile.TemporaryDirectory() as tmp:
            for language, headers in [
                ("CN", ["运单号", "所属包号", "扫描时间", "扫描员"]),
                ("TH", ["หมายเลข AWB", "หมายเลขกระสอบ", "เวลาที่สแกน", "พนักงานสแกน"]),
            ]:
                path = Path(tmp) / (language + ".xlsx")
                fixture(
                    path,
                    headers,
                    [
                        ["old", "", "2026-10-20 14:59:59", "Autopacking_pda01"],
                        ["current", "", "2026-10-20 15:00:00", "Autopacking_pda01"],
                    ],
                )
                inputs = Inputs(
                    {"AUTO_PDA_DATA": str(path)},
                    date(2026, 10, 20),
                    datetime(2026, 10, 21, 0, 5),
                    start_hour=15,
                    log=lambda _: None,
                )
                self.assertEqual(sum(inputs.counter("AUTO_PDA_DATA").values()), 1)


class TemplateTests(unittest.TestCase):
    def test_both_templates_become_values_without_touching_originals(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            paths = input_fixtures(folder)
            for name in ("1AUTOPACKING_BOT.xlsx", "1DWS_&_PDA_v1.8.3 Bot.xlsx"):
                template = ROOT / "Excel" / name
                if not template.exists():
                    self.skipTest("Operator templates are intentionally local-only")
                original = hashlib.sha256(template.read_bytes()).hexdigest()
                output = folder / name
                evidence = calculate(
                    str(template),
                    str(output),
                    date(2026, 10, 20),
                    datetime(2026, 10, 21, 10),
                    paths,
                    log=lambda _: None,
                )
                self.assertGreater(evidence["formulas_replaced"], 1000)
                self.assertEqual(
                    hashlib.sha256(template.read_bytes()).hexdigest(), original
                )
                book = openpyxl.load_workbook(output)
                self.assertFalse(
                    any(c.data_type == "f" for s in book for row in s for c in row)
                )
                if "Autoformat" in book:
                    self.assertEqual(book["Autoformat"]["F4"].value, 201)
                    self.assertEqual(book["Autoformat"]["F18"].value, 199)
                    self.assertEqual(book["Autoformat"]["F37"].value, 2)
                    self.assertEqual(book["Autoformat"]["F38"].value, 0)
                else:
                    self.assertEqual(book["DWSREALTIME"]["F4"].value, 201)
                    self.assertEqual(book["DWSREALTIME"]["F12"].value, 202)
                    self.assertEqual(book["DWSREALTIME"]["F15"].value, 204)
                    self.assertEqual(book["AUTO PDA"]["F5"].value, 203)
                    self.assertEqual(book["REALTIME DB"]["B3"].value, 1)
                book.close()

    def test_failed_and_cancelled_runs_preserve_previous_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            template = folder / "template.xlsx"
            fixture(
                template,
                ["header"],
                [],
                "DWSREALTIME",
                {"AUTO PDA": ([], []), "REALTIME DB": ([], [])},
            )
            book = openpyxl.load_workbook(template)
            book["DWSREALTIME"]["A1"] = "=UNSUPPORTED(1)"
            book.save(template)
            book.close()
            output = folder / "previous.xlsx"
            output.write_bytes(b"previous-complete-artifact")
            with self.assertRaises(UnsupportedFormula):
                calculate(
                    str(template),
                    str(output),
                    date(2026, 10, 20),
                    datetime(2026, 10, 21),
                    {},
                )
            self.assertEqual(output.read_bytes(), b"previous-complete-artifact")
            book = openpyxl.load_workbook(template)
            book["DWSREALTIME"]["A1"] = "=1+2"
            book.save(template)
            book.close()
            with self.assertRaises(InterruptedError):
                calculate(
                    str(template),
                    str(output),
                    date(2026, 10, 20),
                    datetime(2026, 10, 21),
                    {},
                    keep_running=lambda: False,
                )
            self.assertEqual(output.read_bytes(), b"previous-complete-artifact")

    def test_legacy_mode_is_an_explicit_noop(self):
        config = configparser.RawConfigParser()
        config["REPORT_CALCULATION"] = {"enabled": "false"}
        with mock.patch("report_workbooks.calculate") as worker:
            self.assertEqual(
                prepare(
                    config,
                    [SimpleNamespace(key="1", path="unused")],
                    [],
                    "unused",
                    date(2026, 10, 20),
                    datetime(2026, 10, 21),
                    log=lambda _: None,
                ),
                {},
            )
            worker.assert_not_called()

    def test_old_template_raw_cache_is_not_copied_as_fresh_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / "cached-template.xlsx"
            fixture(
                template,
                ["Report"],
                [],
                "Autoformat",
                {"raw_cache": (["Old raw data"], [[i] for i in range(5000)])},
            )
            book = openpyxl.load_workbook(template)
            book["Autoformat"]["A3"] = "=SUM(1,2)"
            book.save(template)
            book.close()
            output = Path(tmp) / "report.xlsx"
            calculate(
                str(template),
                str(output),
                date(2026, 10, 20),
                datetime(2026, 10, 21),
                {},
                log=lambda _: None,
            )
            actual = openpyxl.load_workbook(output)
            self.assertEqual(actual["Autoformat"]["A3"].value, 3)
            self.assertEqual(actual["raw_cache"].max_row, 1)
            actual.close()

    def test_source_failure_happens_before_excel_starts(self):
        import Createphoto

        config = configparser.RawConfigParser()
        config["TIME"] = {"start_hour": "15"}
        workbook = SimpleNamespace(key="1", path="not-opened", display_name="test")
        export = SimpleNamespace(workbook="1")
        with (
            tempfile.TemporaryDirectory() as tmp,
            mock.patch.object(Createphoto, "load_config", return_value=config),
            mock.patch.object(Createphoto, "save_config"),
            mock.patch.object(Createphoto, "migrate_old_export_config"),
            mock.patch.object(Createphoto, "get_workbooks", return_value=[workbook]),
            mock.patch.object(Createphoto, "get_export_items", return_value=[export]),
            mock.patch.object(
                Createphoto, "get_time_config", return_value=config["TIME"]
            ),
            mock.patch(
                "report_workbooks.prepare",
                side_effect=SourceUnavailable("SMB source unavailable"),
            ),
            mock.patch.object(Createphoto.win32, "DispatchEx") as excel,
        ):
            with self.assertRaises(SourceUnavailable):
                Createphoto.run_create(tmp)
            excel.assert_not_called()


if __name__ == "__main__":
    unittest.main()
